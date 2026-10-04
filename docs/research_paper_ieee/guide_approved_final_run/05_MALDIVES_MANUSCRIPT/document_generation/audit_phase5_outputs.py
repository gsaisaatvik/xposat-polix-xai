from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "MALDIVES_Part_01_Canonical_Manuscript.md"
DOCX = ROOT / "MALDIVES_Part_02_IEEE_Manuscript.docx"
TEX = ROOT / "MALDIVES_Part_03_IEEE_Manuscript.tex"
BIB = ROOT / "MALDIVES_Part_04_References.bib"
REPORT = ROOT / "MALDIVES_Part_06_Phase5_QA.md"

TITLE = "An Explainable AI Framework for Analysis of X-Ray Polarimetry Data from XPoSat"
REQUIRED = [
    "I. Introduction", "II. Related Work", "III. Dataset, Scope, and Problem Formulation",
    "IV. Product-Aware Feature Engineering", "V. Explainable Unsupervised Methodology",
    "VI. Independent Harmonic Diagnostic", "VII. Results", "VIII. Discussion",
    "IX. Limitations and Threats to Validity", "X. Conclusion", "Acknowledgment", "References",
]
FORBIDDEN = [
    "POLIX User Handbook", "polix_handbook_2025", "confirmed anomaly", "robust core",
    "all available POLIX data", "novel framework", "median imputation",
]


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def words(text: str) -> int:
    return len(re.findall(r"\b[\w–'-]+\b", text))


md = MD.read_text(encoding="utf-8")
tex = TEX.read_text(encoding="utf-8")
bib = BIB.read_text(encoding="utf-8")
doc = Document(DOCX)
doc_text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    doc_text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)

abstract_match = re.search(r"(?s)## Abstract\s*(.*?)\s*\*\*Index Terms", md)
abstract_words = words(abstract_match.group(1)) if abstract_match else 0
cite_keys = set(re.findall(r"@([A-Za-z0-9_:-]+)", md))
bib_keys = set(re.findall(r"@[A-Za-z]+\{([^,]+),", bib))

checks = []
def check(name, passed, detail):
    checks.append((name, bool(passed), detail))

check("Exact title in Markdown", TITLE in md, TITLE)
check("Exact title in DOCX", TITLE in doc_text, TITLE)
check("Exact title in LaTeX", TITLE in tex, TITLE)
check("Abstract length", 190 <= abstract_words <= 230, f"{abstract_words} words")
check("Required Markdown sections", all(x in md for x in REQUIRED), f"{len(REQUIRED)} required markers")
check("Four figures in DOCX", len(doc.inline_shapes) == 4, f"{len(doc.inline_shapes)} inline figures")
check("Four tables in DOCX", len(doc.tables) == 4, f"{len(doc.tables)} tables")
check("DOCX section geometry", len(doc.sections) >= 2, f"{len(doc.sections)} Word sections")
check("Citation keys resolve", cite_keys <= bib_keys, f"missing={sorted(cite_keys-bib_keys)}")
check("No uncited BibTeX entries", bib_keys <= cite_keys, f"uncited={sorted(bib_keys-cite_keys)}")
check("Seventeen references", len(bib_keys) == 17, f"{len(bib_keys)} BibTeX entries")
check("No forbidden phrases", not any(x.lower() in (md+tex+bib).lower() for x in FORBIDDEN), "checked Markdown, LaTeX, BibTeX")
check("Fixed result present", "21 observations Normal and flagged four candidates" in md, "21 Normal / four candidates")
check("Stable-three wording present", "three-candidate Matrix-C core stable under the tested procedures" in md, "procedure-qualified")
check("Sco X-1 ranking present", all(x in md for x in ["energy peak channel (2.454)", "energy weighted mean channel (2.207)", "energy-channel entropy (1.814)"]), "2.454 / 2.207 / 1.814")
check("Blank-sky rule present", "Thirteen of the 15" in md and "reduced chi-square at most 2" in md, "13 of 15, reduced chi-square <= 2")
check("No final PDF in phase folder", not any(ROOT.glob("*.pdf")), "temporary QA files excluded")

xml_ok = True
with zipfile.ZipFile(DOCX) as zf:
    for name in zf.namelist():
        if name.endswith(".xml"):
            try:
                ET.fromstring(zf.read(name))
            except ET.ParseError:
                xml_ok = False
                break
check("DOCX package XML is well formed", xml_ok, "all XML parts parsed")

status = "PASS" if all(p for _, p, _ in checks) else "FAIL"
lines = [
    "# MALDIVES Part 06 — Phase-5 Structural QA",
    "",
    f"Overall automated status: **{status}**",
    "",
    "| Check | Status | Detail |",
    "|---|---|---|",
]
for name, passed, detail in checks:
    lines.append(f"| {name} | {'PASS' if passed else 'FAIL'} | {detail} |")
lines.extend([
    "",
    "## File hashes",
    "",
    "| File | SHA-256 |",
    "|---|---|",
])
for path in [MD, DOCX, TEX, BIB, ROOT / "MALDIVES_Part_05_Claim_Evidence_Ledger.md"]:
    lines.append(f"| `{path.name}` | `{sha(path)}` |")
lines.extend([
    "",
    "## Visual-render status",
    "",
    "The required Word visual-render attempt was performed twice. The bundled renderer could not locate LibreOffice, and the local Microsoft Word conversion did not complete within the bounded QA windows. No rendered-page pass is claimed. The DOCX package, section structure, figures, tables, citations, required claims, and XML were audited directly. Phase 6 must repeat page-image inspection on a working Word/LibreOffice renderer before any submission-ready status is assigned.",
])
REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"STATUS={status}")
print(f"ABSTRACT_WORDS={abstract_words}")
print(f"DOCX_PARAGRAPHS={len(doc.paragraphs)}")
print(f"DOCX_TABLES={len(doc.tables)}")
print(f"DOCX_INLINE_SHAPES={len(doc.inline_shapes)}")
print(f"DOCX_SECTIONS={len(doc.sections)}")
print(f"WROTE={REPORT}")
