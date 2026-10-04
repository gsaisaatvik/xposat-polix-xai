from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "SINGAPORE_Part_00_Controlled_Manuscript.md"
OUTPUT = ROOT / "SINGAPORE_Part_01_Guide_Ready_Manuscript.docx"
BUILD_OUTPUT = ROOT / "qa_temp" / "SINGAPORE_Guide_Ready_build.docx"

CITATION_ORDER = [
    "isro_xposat", "rishin2010thomson", "fabiani2018instrumentation",
    "kislat2015stokes", "issdc_xposat_archive", "baron2017weirdest",
    "giles2019serendipity", "lochner2021astronomaly",
    "li2024explainable_anomaly", "lundberg2017shap", "yeh2019infidelity",
    "campos2016evaluation", "pearson1901pca", "macqueen1967kmeans",
    "liu2008isolation", "pedregosa2011sklearn", "issdc_xposat_ack",
]
CITE_NUM = {key: i + 1 for i, key in enumerate(CITATION_ORDER)}

FIGURE_CAPTIONS = {
    1: "Product-aware screening and independent harmonic-diagnostic architecture. PCA, KMeans, and Isolation Forest receive the same standardized Matrix-C rows; only Isolation Forest defines the fixed label. WeightedRoll is analysed separately, and the branches are compared without confidence fusion.",
    2: "Two-component PCA projection of the standardized Matrix-C observations. Circles denote source observations, squares denote blank skies, and open diamonds outline the fixed four candidates. The projection is descriptive and does not determine or validate the label.",
    3: "Fixed Isolation Forest anomaly-score ranking for all 25 observations. Hatched bars identify the fixed candidates; annotations report selection counts across 100 tested seeds. The counts are algorithmic frequencies rather than confidence levels.",
    4: "Fractional second-harmonic coordinates for all 25 observations. Filled squares form the 13-fit blank-sky reference; open squares are excluded by the fit rule; open diamonds identify the fixed ML candidates. The cross is the component-wise sample mean plus or minus one sample standard deviation, not a confidence or detection contour.",
}


def parse_blocks(text: str):
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("$$"):
            equation = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("$$"):
                equation.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(("equation", "\n".join(equation).strip()))
            continue
        if re.match(r"^!\[.*\]\(.*\)$", line):
            match = re.match(r"^!\[(.*)\]\((.*)\)$", line)
            blocks.append(("image", match.group(1), match.group(2)))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[|:\-\s]+\|$", lines[i + 1].strip()):
            rows = [[c.strip() for c in line.strip().strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
            continue
        if line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
            i += 1
            continue
        if line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            blocks.append(("list", re.sub(r"^\d+\.\s+", "", line)))
            i += 1
            continue
        para = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if (not nxt or nxt.startswith("#") or nxt.startswith("$$") or nxt.startswith("|")
                    or re.match(r"^!\[.*\]\(.*\)$", nxt) or re.match(r"^\d+\.\s+", nxt)):
                break
            para.append(nxt)
            i += 1
        blocks.append(("paragraph", " ".join(para)))
    return blocks


def replace_citations(text: str) -> str:
    def repl(match):
        keys = [k.strip().lstrip("@") for k in match.group(1).split(";")]
        return ", ".join(f"[{CITE_NUM[k]}]" for k in keys)
    return re.sub(r"\[@([^\]]+)\]", repl, text)


def set_font(run, name="Times New Roman", size=10.5, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def plain_math(text: str) -> str:
    """Render inline LaTeX-like notation as portable, readable Unicode text."""
    replacements = [
        (r"\mathbb{R}", "ℝ"),
        (r"\times", "×"),
        (r"\in", "∈"),
        (r"\ge", "≥"),
        (r"\le", "≤"),
        (r"\ldots", "…"),
        (r"\phi", "φ"),
        (r"\sigma", "σ"),
        (r"\chi", "χ"),
        (r"\nu", "ν"),
        (r"\rho", "ρ"),
        (r"\tfrac12", "½"),
        (r"\leftarrow", "←"),
        (r"\bmod", "mod"),
        (r"^\circ", "°"),
        (r"\_", "_"),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
    text = re.sub(r"\\(?:mathrm|operatorname|text)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\bar\s*([A-Za-z])", lambda match: match.group(1) + "\u0304", text)
    text = re.sub(r"_\{([^{}]*)\}", r"_(\1)", text)
    text = re.sub(r"\^\{([^{}]*)\}", r"^(\1)", text)
    return text.replace("{", "").replace("}", "").replace("\\", "")


def add_inline(paragraph, text: str, size=10.5):
    text = replace_citations(text)
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|\$.*?\$|\*.*?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=size, bold=True)
        elif token.startswith("`"):
            set_font(paragraph.add_run(token[1:-1]), name="Consolas", size=max(8.0, size - 1.0))
        elif token.startswith("$"):
            set_font(paragraph.add_run(plain_math(token[1:-1])), name="Cambria Math", size=size, italic=True)
        else:
            set_font(paragraph.add_run(token[1:-1]), size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]), size=size)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_shading(cell, fill="E8EEF5"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_table_geometry(table, widths_in):
    widths = [int(round(x * 1440)) for x in widths_in]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[index])
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = p.add_run()
    run._r.append(fld)


def build():
    blocks = parse_blocks(MASTER.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4)

    h1 = styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h1.font.size = Pt(11.5)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(5)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h2.font.size = Pt(10.5)
    h2.font.bold = True
    h2.font.italic = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(7)
    h2.paragraph_format.space_after = Pt(3)
    h2.paragraph_format.keep_with_next = True

    if "Equation" not in [s.name for s in styles]:
        eq_style = styles.add_style("Equation", 1)
    else:
        eq_style = styles["Equation"]
    eq_style.font.name = "Cambria Math"
    eq_style.font.size = Pt(10.5)
    eq_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq_style.paragraph_format.space_before = Pt(4)
    eq_style.paragraph_format.space_after = Pt(4)
    eq_style.paragraph_format.keep_together = True

    cap_style = styles["Caption"]
    cap_style.font.name = "Times New Roman"
    cap_style.font.size = Pt(9)
    cap_style.font.italic = False
    cap_style.font.bold = False
    cap_style.font.color.rgb = RGBColor(0, 0, 0)
    cap_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cap_style.paragraph_format.space_after = Pt(6)
    cap_style.paragraph_format.keep_together = True

    title_seen = False
    front_matter = True
    abstract_mode = False
    references_mode = False
    pending_table_title = None
    figure_number = 0
    equation_number = 0
    table_number = 0

    widths = {
        1: [1.35, 2.10, 3.05],
        2: [1.30, 0.55, 4.65],
        3: [1.15, 0.75, 0.75, 1.00, 2.85],
        4: [1.05, 0.80, 0.85, 0.85, 0.85, 2.10],
    }

    for block in blocks:
        kind = block[0]
        if kind == "h1" and not title_seen:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(block[1])
            set_font(run, size=17, bold=True)
            title_seen = True
            continue
        if kind == "h2" and block[1] == "Abstract":
            abstract_mode = True
            front_matter = False
            continue
        if kind == "h2" and block[1] == "References":
            references_mode = True
            doc.add_paragraph("REFERENCES", style="Heading 1")
            continue
        if kind == "h2" and block[1] == "Acknowledgment":
            doc.add_paragraph("ACKNOWLEDGMENT", style="Heading 1")
            continue
        if kind == "h1":
            doc.add_paragraph(block[1].upper(), style="Heading 1")
            continue
        if kind == "h2":
            doc.add_paragraph(block[1], style="Heading 2")
            continue
        if kind == "paragraph":
            text = block[1]
            if text.startswith("**TABLE "):
                pending_table_title = re.sub(r"\*\*", "", text)
                continue
            if text.startswith("**Fig. "):
                continue
            p = doc.add_paragraph()
            if not abstract_mode and not references_mode and not front_matter:
                p.paragraph_format.first_line_indent = Inches(0.2)
            if front_matter:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, text, size=10)
            elif abstract_mode and text.startswith("**Index Terms"):
                add_inline(p, text, size=9.5)
                p.paragraph_format.space_after = Pt(8)
                abstract_mode = False
            elif abstract_mode:
                label = p.add_run("Abstract—")
                set_font(label, size=9.5, bold=True)
                add_inline(p, text, size=9.5)
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.right_indent = Inches(0.25)
            elif references_mode:
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, text, size=9)
            else:
                add_inline(p, text)
            continue
        if kind == "list":
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, block[1])
            continue
        if kind == "equation":
            equation_number += 1
            p = doc.add_paragraph(style="Equation")
            p.add_run(f"[[EQ{equation_number:02d}]]")
            continue
        if kind == "table":
            table_number += 1
            if pending_table_title:
                p = doc.add_paragraph(style="Caption")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                run = p.add_run(pending_table_title.upper())
                set_font(run, size=9, bold=True)
                pending_table_title = None
            rows = block[1]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            set_table_geometry(table, widths[table_number])
            set_repeat_table_header(table.rows[0])
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell)
                    if ri == 0:
                        set_shading(cell)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    if ri == 0 or (ci > 0 and len(value) < 18):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    add_inline(p, value, size=8.6)
                    if ri == 0:
                        for run in p.runs:
                            run.bold = True
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(2)
            continue
        if kind == "image":
            figure_number += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            figure_width = 5.55 if figure_number == 1 else 6.35
            shape = p.add_run().add_picture(block[2], width=Inches(figure_width))
            shape._inline.docPr.set("descr", FIGURE_CAPTIONS[figure_number])
            shape._inline.docPr.set("title", f"Figure {figure_number}")
            cp = doc.add_paragraph(style="Caption")
            label = cp.add_run(f"Fig. {figure_number}. ")
            set_font(label, size=9, bold=False, italic=True)
            set_font(cp.add_run(FIGURE_CAPTIONS[figure_number]), size=9, bold=False)
            continue

    core = doc.core_properties
    core.title = "An Explainable AI Framework for Analysis of X-Ray Polarimetry Data from XPoSat"
    core.subject = "Guide-review manuscript; venue adaptation pending"
    core.author = "[Student authors pending guide approval]"
    core.keywords = "XPoSat; POLIX; explainable AI; unsupervised anomaly screening"
    core.comments = "Single-column guide-review copy. Final IEEE venue formatting remains pending."
    BUILD_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(BUILD_OUTPUT)
    BUILD_OUTPUT.replace(OUTPUT)
    print(f"WROTE={OUTPUT}")
    print(f"EQUATION_PLACEHOLDERS={equation_number}")
    print(f"FIGURES={figure_number}")
    print(f"TABLES={table_number}")


if __name__ == "__main__":
    build()
