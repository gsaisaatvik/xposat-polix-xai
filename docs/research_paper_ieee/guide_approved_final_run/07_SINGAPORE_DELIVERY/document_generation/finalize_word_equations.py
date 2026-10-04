from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PHASE_ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = PHASE_ROOT / "SINGAPORE_Part_01_Guide_Ready_Manuscript.docx"

EQUATIONS = {
    "[[EQ01]]": "j_peak = arg max_j n_j,    j̄ = (Σ_j j n_j)/(Σ_j n_j),    s_j = √[(Σ_j (j−j̄)² n_j)/(Σ_j n_j)]",
    "[[EQ02]]": "P_j = |x_j l_1j| r_1 + |x_j l_2j| r_2,    K_j = (x_j − c_j)²",
    "[[EQ03]]": "I_j = s(x) − s(x^(j→0)),    Z_j = |x_j|",
    "[[EQ04]]": "E_j = N(P)_j + N(K)_j + N(max(I, 0))_j + N(Z)_j",
    "[[EQ05]]": "χ²_red = (1/ν) Σ_i [(y_i − y_i^(mod))/σ_i]²",
    "[[EQ06]]": "A = √(Q² + U²),    m_raw = 100 A/C",
    "[[EQ07]]": "φ_fit = 0.5 atan2(U, Q) mod 180°,    q = Q/C,    u = U/C",
    "[[EQ08]]": "z_m = (m_raw − m̄_blank)/s_blank",
}


def make_omath(text: str):
    math_para = OxmlElement("m:oMathPara")
    math_para_pr = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "centerGroup")
    math_para_pr.append(justification)
    math_para.append(math_para_pr)

    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_run_pr = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    math_run_pr.append(style)
    math_run.append(math_run_pr)
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    return math_para


def main():
    document = Document(DOCX_PATH)
    replaced = 0
    for paragraph in document.paragraphs:
        marker = paragraph.text.strip()
        if marker not in EQUATIONS:
            continue
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)
        paragraph._p.append(make_omath(EQUATIONS[marker]))
        replaced += 1
    if replaced != len(EQUATIONS):
        raise RuntimeError(f"Expected {len(EQUATIONS)} equation placeholders; replaced {replaced}.")
    document.save(DOCX_PATH)
    print(f"NATIVE_EQUATIONS={replaced}")
    print(f"WROTE={DOCX_PATH}")


if __name__ == "__main__":
    main()
