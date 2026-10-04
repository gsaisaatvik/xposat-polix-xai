from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import statistics
from pathlib import Path

from docx import Document


WEBAPP = Path(r"D:\polix_xai_webapp")
ARCHIVE = Path(r"D:\ISROtrial\Polix_L2_full_archive")
RUN = WEBAPP / "research_paper_ieee" / "guide_approved_final_run"
PHASE5 = RUN / "05_MALDIVES_MANUSCRIPT"
OUT = RUN / "06_AFGHANISTAN_REVIEW" / "working" / "phase6_checks.json"
SUPP = WEBAPP / "research_paper_ieee" / "supplementary_experiments"
PHYS = ARCHIVE / "final_project_outputs" / "04_polarimetry_results"


def read_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def approx(a, b, tol=5e-7):
    return abs(float(a) - float(b)) <= tol


checks = []


def add(group, claim, passed, actual, expected, source):
    checks.append({
        "group": group,
        "claim": claim,
        "passed": bool(passed),
        "actual": actual,
        "expected": expected,
        "source": str(source),
    })


# Matrix C
matrix_path = ARCHIVE / "polix_matrix_v2_C_primary_plus_supporting.csv"
matrix = read_csv(matrix_path)
features = [k for k in matrix[0] if k != "observation_id"]
missing = sum(v is None or str(v).strip() == "" for row in matrix for k, v in row.items() if k != "observation_id")
add("dataset", "Matrix C has 25 rows", len(matrix) == 25, len(matrix), 25, matrix_path)
add("dataset", "Matrix C has 15 features", len(features) == 15, len(features), 15, matrix_path)
add("dataset", "Matrix C has zero missing feature values", missing == 0, missing, 0, matrix_path)
add("dataset", "WeightedRoll is absent from Matrix C", not any("weightedroll" in x.lower() for x in features), features, "no WeightedRoll feature", matrix_path)

truth_path = RUN / "01_NEPAL_EVIDENCE" / "NEPAL_Part_03_All_25_Observation_Truth_Table.csv"
truth = read_csv(truth_path)
role_counts = {
    "source": sum(r["observation_role"] == "source" for r in truth),
    "blank_sky": sum(r["observation_role"] == "blank_sky" for r in truth),
}
add("dataset", "Project role mapping is 10 source and 15 blank sky", role_counts == {"source": 10, "blank_sky": 15}, role_counts, {"source": 10, "blank_sky": 15}, truth_path)

# Fixed reproduction
fixed_path = SUPP / "deployed_model_reproduction.csv"
fixed = read_csv(fixed_path)
anomalies = [r for r in fixed if r["recomputed_prediction"] == "Anomaly"]
expected_anoms = {
    "X01_PLX_C24_0018_000000", "X01_PLX_G01_0006_000000",
    "X01_PLX_G01_0003_000000", "X01_PLX_C24_0010_000000",
}
add("fixed", "Fixed result has four candidates", len(anomalies) == 4, len(anomalies), 4, fixed_path)
add("fixed", "Fixed result has 21 Normal rows", sum(r["recomputed_prediction"] == "Normal" for r in fixed) == 21, sum(r["recomputed_prediction"] == "Normal" for r in fixed), 21, fixed_path)
add("fixed", "Fixed candidate identities match", {r["observation_id"] for r in anomalies} == expected_anoms, sorted(r["observation_id"] for r in anomalies), sorted(expected_anoms), fixed_path)
add("fixed", "All stored and recomputed labels match", all(r["prediction_match"] == "True" for r in fixed), sum(r["prediction_match"] == "True" for r in fixed), 25, fixed_path)
add("fixed", "Maximum score replay difference is zero", max(float(r["absolute_score_difference"]) for r in fixed) == 0.0, max(float(r["absolute_score_difference"]) for r in fixed), 0.0, fixed_path)

expected_scores = {
    "X01_PLX_C24_0018_000000": 0.615964,
    "X01_PLX_G01_0006_000000": 0.593516,
    "X01_PLX_G01_0003_000000": 0.572308,
    "X01_PLX_C24_0010_000000": 0.517611,
}
actual_scores = {r["observation_id"]: float(r["recomputed_anomaly_score"]) for r in anomalies}
add("fixed", "Fixed candidate scores match manuscript rounding", all(round(actual_scores[k], 6) == v for k, v in expected_scores.items()), {k: round(v, 6) for k, v in actual_scores.items()}, expected_scores, fixed_path)

# Stability
seed_path = SUPP / "isolation_seed_stability_summary.csv"
seeds = {r["observation_id"]: int(r["times_flagged"]) for r in read_csv(seed_path)}
expected_seed = {
    "X01_PLX_C24_0018_000000": 100,
    "X01_PLX_G01_0006_000000": 100,
    "X01_PLX_G01_0003_000000": 100,
    "X01_PLX_C24_0010_000000": 29,
    "X01_PLX_C24_0020_000000": 68,
    "X01_PLX_P01_0005_000000": 3,
}
add("stability", "Seed-selection counts match", all(seeds[k] == v for k, v in expected_seed.items()), {k: seeds[k] for k in expected_seed}, expected_seed, seed_path)

contam_path = SUPP / "contamination_candidate_stability.csv"
contam = {r["observation_id"]: int(r["settings_flagged"]) for r in read_csv(contam_path)}
expected_contam = {
    "X01_PLX_C24_0018_000000": 4,
    "X01_PLX_G01_0006_000000": 4,
    "X01_PLX_G01_0003_000000": 4,
    "X01_PLX_C24_0010_000000": 3,
    "X01_PLX_C24_0020_000000": 2,
    "X01_PLX_P01_0005_000000": 1,
}
add("stability", "Contamination-setting persistence matches", all(contam[k] == v for k, v in expected_contam.items()), {k: contam[k] for k in expected_contam}, expected_contam, contam_path)

jack_path = SUPP / "jackknife_observation_stability.csv"
jack = {r["observation_id"]: int(r["times_flagged_when_in_training"]) for r in read_csv(jack_path)}
expected_jack = {
    "X01_PLX_C24_0018_000000": 24,
    "X01_PLX_G01_0006_000000": 24,
    "X01_PLX_G01_0003_000000": 24,
    "X01_PLX_C24_0010_000000": 19,
    "X01_PLX_C24_0020_000000": 7,
    "X01_PLX_P01_0005_000000": 2,
}
add("stability", "Included-observation jackknife counts match", all(jack[k] == v for k, v in expected_jack.items()), {k: jack[k] for k in expected_jack}, expected_jack, jack_path)

jack_runs_path = SUPP / "jackknife_run_summary.csv"
rhos = sorted(float(r["spearman_rho_vs_full_on_common_observations"]) for r in read_csv(jack_runs_path))
add("stability", "Jackknife median rank correlation matches", approx(statistics.median(rhos), 0.989565217391304, 1e-12), statistics.median(rhos), 0.989565217391304, jack_runs_path)
add("stability", "Jackknife minimum rank correlation matches", approx(min(rhos), 0.968695652173913, 1e-12), min(rhos), 0.968695652173913, jack_runs_path)

# Matrix tiers
ablation_path = SUPP / "matrix_ablation_summary.csv"
ablation = {r["matrix"]: set(x.strip() for x in r["isolation_anomaly_ids"].split("|") if x.strip()) for r in read_csv(ablation_path)}
persistent = set.intersection(ablation["A"], ablation["B"], ablation["C"])
expected_persistent = {"X01_PLX_C24_0018_000000", "X01_PLX_G01_0006_000000"}
add("ablation", "Cross-tier persistent pair matches", persistent == expected_persistent, sorted(persistent), sorted(expected_persistent), ablation_path)

# Rank correlations
rank_path = SUPP / "ranking_agreement.csv"
rank_rows = read_csv(rank_path)
rank_actual = {(r["metric_left"], r["metric_right"]): float(r["spearman_rho"]) for r in rank_rows}
rank_expected = {
    ("PCA distance", "Isolation Forest anomaly score"): 0.8946153846153846,
    ("PCA distance", "KMeans centroid distance"): 0.15579919503434767,
    ("KMeans centroid distance", "Isolation Forest anomaly score"): 0.18657434467076198,
}
add(
    "ranking",
    "Spearman ranking correlations match",
    all(approx(rank_actual[k], v, 1e-12) for k, v in rank_expected.items()),
    {" vs ".join(k): v for k, v in rank_actual.items()},
    {" vs ".join(k): v for k, v in rank_expected.items()},
    rank_path,
)

# XAI
xai_path = SUPP / "deployed_xai_exact_from_model_service.csv"
xai = sorted(read_csv(xai_path), key=lambda r: int(r["exact_rank"]))
top3 = [(r["feature"], float(r["xai_score"])) for r in xai[:3]]
expected_top3 = [
    ("t1A_energy_peak_channel", 2.4544990315347768),
    ("t1A_energy_weighted_mean_channel", 2.207369),
    ("t1A_energy_channel_entropy", 1.813968),
]
add("xai", "Sco X-1 feature order matches", [x[0] for x in top3] == [x[0] for x in expected_top3], top3, expected_top3, xai_path)
add("xai", "Sco X-1 main-paper rounding matches", [round(x[1], 3) for x in top3] == [2.454, 2.207, 1.814], [round(x[1], 3) for x in top3], [2.454, 2.207, 1.814], xai_path)

faith_path = SUPP / "faithfulness_verdict_reproduction.csv"
faith = read_csv(faith_path)
faith_counts = {name: sum(r["overall_xai_faithfulness"] == name for r in faith) for name in ("Strong", "Moderate", "Weak")}
add("xai", "Six-case verdict count matches", faith_counts == {"Strong": 5, "Moderate": 1, "Weak": 0}, faith_counts, {"Strong": 5, "Moderate": 1, "Weak": 0}, faith_path)

# Harmonic branch
harm_path = PHYS / "path2_final_polarization_xai_result_table_with_roles.csv"
harm = read_csv(harm_path)
quality_counts = {
    "acceptable": sum(r["fit_quality"] == "acceptable" for r in harm),
    "caution": sum(r["fit_quality"] == "caution" for r in harm),
    "poor": sum(r["fit_quality"] == "poor_simple_sinusoid_fit" for r in harm),
}
add("harmonic", "All 25 harmonic rows are present", len(harm) == 25, len(harm), 25, harm_path)
add("harmonic", "Fit-quality counts match", quality_counts == {"acceptable": 19, "caution": 3, "poor": 3}, quality_counts, {"acceptable": 19, "caution": 3, "poor": 3}, harm_path)

baseline_path = PHYS / "path2_blank_sky_modulation_baseline.csv"
baseline_rows = read_csv(baseline_path)
baseline = next(r for r in baseline_rows if r["baseline_set"] == "acceptable_blank_sky_only")
add("harmonic", "Selected blank-sky reference has 13 observations", int(baseline["n_observations"]) == 13, int(baseline["n_observations"]), 13, baseline_path)
add("harmonic", "Blank-sky mean raw modulation matches", approx(baseline["mean_raw_modulation_percent"], 1.147820338848039, 1e-12), float(baseline["mean_raw_modulation_percent"]), 1.147820338848039, baseline_path)
add("harmonic", "Blank-sky sample spread matches", approx(baseline["std_raw_modulation_percent"], 0.565959672908476, 1e-12), float(baseline["std_raw_modulation_percent"]), 0.565959672908476, baseline_path)

source_path = PHYS / "path2_source_vs_blank_sky_modulation_comparison.csv"
sources = read_csv(source_path)
add("harmonic", "Ten source comparison rows are present", len(sources) == 10, len(sources), 10, source_path)
add("harmonic", "All source rows satisfy the declared empirical rule", all(r["source_modulation_vs_blank_status"] == "within_blank_sky_baseline_range" for r in sources), sum(r["source_modulation_vs_blank_status"] == "within_blank_sky_baseline_range" for r in sources), 10, source_path)

representative_expected = {
    "X01_PLX_G01_0006_000000": (1.1395784576346573, 2.0307608047722137, "caution"),
    "X01_PLX_G01_0003_000000": (0.596631928790355, 57.43351375364045, "poor_simple_sinusoid_fit"),
    "X01_PLX_P01_0005_000000": (1.6973459072645685, 1.0625397079395509, "acceptable"),
    "X01_PLX_C24_0018_000000": (0.8539506414795067, 2.565811346028652, "caution"),
    "X01_PLX_C24_0010_000000": (1.51853532343414, 0.7134283152346765, "acceptable"),
}
harm_by_id = {r["observation_id"]: r for r in harm}
rep_ok = all(
    approx(harm_by_id[k]["raw_modulation_percent"], v[0], 1e-12)
    and approx(harm_by_id[k]["reduced_chi2"], v[1], 1e-12)
    and harm_by_id[k]["fit_quality"] == v[2]
    for k, v in representative_expected.items()
)
add("harmonic", "Representative table values match", rep_ok, {k: (harm_by_id[k]["raw_modulation_percent"], harm_by_id[k]["reduced_chi2"], harm_by_id[k]["fit_quality"]) for k in representative_expected}, representative_expected, harm_path)

# Cross-format content checks
md_path = PHASE5 / "MALDIVES_Part_01_Canonical_Manuscript.md"
tex_path = PHASE5 / "MALDIVES_Part_03_IEEE_Manuscript.tex"
docx_path = PHASE5 / "MALDIVES_Part_02_IEEE_Manuscript.docx"
bib_path = PHASE5 / "MALDIVES_Part_04_References.bib"
md = md_path.read_text(encoding="utf-8")
tex = tex_path.read_text(encoding="utf-8")
bib = bib_path.read_text(encoding="utf-8")
doc = Document(docx_path)
doc_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
title = "An Explainable AI Framework for Analysis of X-Ray Polarimetry Data from XPoSat"
key_tokens = ["21 observations Normal", "Blank Sky-13", "Sco X-1", "Her X-1", "Blank Sky-5", "2.454", "2.207", "1.814", "1.147820", "0.565960", "57.433514", "100 tested seeds"]
add("formats", "Exact title appears in all three formats", all(title in x for x in (md, tex, doc_text)), True, True, "Markdown/DOCX/LaTeX")
add("formats", "Key numerical tokens agree across formats", all(all(token in x for x in (md, tex, doc_text)) for token in key_tokens), [t for t in key_tokens if not all(t in x for x in (md, tex, doc_text))], [], "Markdown/DOCX/LaTeX")
add("formats", "DOCX contains four figures and four tables", len(doc.inline_shapes) == 4 and len(doc.tables) == 4, {"figures": len(doc.inline_shapes), "tables": len(doc.tables)}, {"figures": 4, "tables": 4}, docx_path)
cite_keys = set(re.findall(r"@([A-Za-z0-9_:-]+)", md))
bib_keys = set(re.findall(r"@[A-Za-z]+\{([^,]+),", bib))
add("formats", "Citation and BibTeX keys are one-to-one", cite_keys == bib_keys and len(bib_keys) == 17, {"cited": sorted(cite_keys), "bib": sorted(bib_keys)}, "17 matching keys", bib_path)

forbidden = ["polix handbook", "polix_handbook", "confirmed anomaly", "robust core", "median imputation", "all available POLIX data", "novel framework"]
found_forbidden = [x for x in forbidden if x in (md + tex + bib).lower()]
add("language", "Forbidden or superseded wording is absent", not found_forbidden, found_forbidden, [], "Markdown/LaTeX/BibTeX")

# Protected artifacts
manifest_path = RUN / "00_INDIA_CONTROL" / "INDIA_Part_03_Protected_File_Hashes.csv"
manifest = read_csv(manifest_path)
mismatches = []
for row in manifest:
    path = Path(row["path"])
    if not path.exists():
        mismatches.append((str(path), "MISSING"))
        continue
    h = hashlib.sha256(path.read_bytes()).hexdigest().upper()
    if h != row["sha256"].upper():
        mismatches.append((str(path), h))
add("integrity", "All protected artifacts are unchanged", not mismatches, mismatches, [], manifest_path)

OUT.write_text(json.dumps({
    "summary": {
        "checks": len(checks),
        "passed": sum(c["passed"] for c in checks),
        "failed": sum(not c["passed"] for c in checks),
    },
    "checks": checks,
}, indent=2, ensure_ascii=False), encoding="utf-8")
print(json.dumps({"checks": len(checks), "passed": sum(c["passed"] for c in checks), "failed": sum(not c["passed"] for c in checks)}, indent=2))
for c in checks:
    if not c["passed"]:
        print("FAIL", c["group"], c["claim"], c["actual"])
