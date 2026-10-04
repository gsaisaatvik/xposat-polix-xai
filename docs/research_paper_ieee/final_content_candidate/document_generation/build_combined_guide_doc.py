from pathlib import Path
import re
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"D:\polix_xai_webapp\research_paper_ieee\final_content_candidate")
OUT = ROOT / "POLIX_Maam_Guide_Review_Sample_Draft.docx"
GUIDE = ROOT / "14_GUIDE_REVIEW_PACKAGE.md"
DECISIONS = ROOT / "09_FINAL_GUIDE_DECISION_SHEET.md"
PAPER = ROOT / "01_FINAL_MANUSCRIPT_DRAFT2.md"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
NAVY = RGBColor(25, 48, 73)
GRAY = RGBColor(90, 98, 108)

def font(run, size=11, bold=None, italic=None, color=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color

def clean(s):
    s = re.sub(r"\s*\[@[^\]]+\]", "", s)
    s = s.replace("The acknowledgment wording follows the current ISSDC guidance.", "")
    s = s.replace("Final human citation, number, and prose verification.", "Final human verification of numbers and prose.")
    s = s.replace(", controlled references", "").replace(", citation recheck", "")
    s = s.replace(chr(96), "").replace("**", "").replace("*", "")
    for a,b in {"—":" - ","–":"-","≤":"<=","±":"+/-","μ":"mu","χ":"chi","ν":"nu","φ":"phi","ψ":"psi",
                "\\le":"<=","\\pm":"+/-","\\mu":"mu","\\chi":"chi","\\phi":"phi","\\psi":"psi",
                "\\bar q":"q_bar","\\bar u":"u_bar","\\sum":"sum","\\sqrt":"sqrt","\\cos":"cos",
                "\\sin":"sin","\\max":"max","\\operatorname":"","\\mathbf":"","\\mathrm":"","\\text":""}.items():
        s=s.replace(a,b)
    s=s.replace("\\(","").replace("\\)","").replace("\\[","").replace("\\]","")
    s=re.sub(r"\\([A-Za-z]+)",r"\1",s).replace("{","(").replace("}",")")
    return re.sub(r"\s+"," ",s).strip()

def para(doc, text, style=None, align=None, italic=False):
    p=doc.add_paragraph(style=style)
    if align is not None: p.alignment=align
    r=p.add_run(clean(text)); font(r,italic=italic)
    return p

def shade(cell, fill):
    tcpr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcpr.append(shd)

def cell_margins(cell):
    tcpr=cell._tc.get_or_add_tcPr(); mar=OxmlElement("w:tcMar")
    for side,val in (("top",80),("start",120),("bottom",80),("end",120)):
        e=OxmlElement("w:"+side); e.set(qn("w:w"),str(val)); e.set(qn("w:type"),"dxa"); mar.append(e)
    tcpr.append(mar)

def table_widths(headers):
    key=" ".join(headers).lower(); n=len(headers)
    if n==5 and "seed" in key: return [1900,1000,1100,1100,4260]
    if n==5: return [1800,1100,1600,1100,3760]
    if n==4 and "decision" in key: return [600,1800,4800,2160]
    if n==4: return [1500,2200,2700,2960]
    if n==3: return [1700,900,6760]
    if n==2: return [2700,6660]
    return [9360//n]*n

def geometry(table,widths):
    table.autofit=False; pr=table._tbl.tblPr
    for tag,val,typ in (("tblW",sum(widths),"dxa"),("tblInd",120,"dxa")):
        e=OxmlElement("w:"+tag); e.set(qn("w:w"),str(val)); e.set(qn("w:type"),typ); pr.append(e)
    lay=OxmlElement("w:tblLayout"); lay.set(qn("w:type"),"fixed"); pr.append(lay)
    grid=table._tbl.tblGrid
    for x in list(grid): grid.remove(x)
    for w in widths:
        e=OxmlElement("w:gridCol"); e.set(qn("w:w"),str(w)); grid.append(e)
    for row in table.rows:
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]/1440); c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER; cell_margins(c)
            tcw=c._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is None: tcw=OxmlElement("w:tcW"); c._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn("w:w"),str(widths[j])); tcw.set(qn("w:type"),"dxa")

def add_table(doc,rows):
    if not rows:return
    n=len(rows[0]); rows=[r[:n]+[""]*max(0,n-len(r)) for r in rows]
    t=doc.add_table(rows=1,cols=n); t.style="Table Grid"; widths=table_widths(rows[0])
    for j,v in enumerate(rows[0]):
        c=t.rows[0].cells[j]; c.text=clean(v); shade(c,"F2F4F7")
        for r in c.paragraphs[0].runs: font(r,9,bold=True,color=NAVY)
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    trpr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); trpr.append(rep)
    for vals in rows[1:]:
        cells=t.add_row().cells
        for j,v in enumerate(vals):
            cells[j].text=clean(v)
            for r in cells[j].paragraphs[0].runs: font(r,9)
            cells[j].paragraphs[0].paragraph_format.space_after=Pt(0)
            cells[j].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER if len(v)<28 else WD_ALIGN_PARAGRAPH.LEFT
    geometry(t,widths); doc.add_paragraph()

def add_list(doc,text,numbered=False):
    p=doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.left_indent=Inches(.5); p.paragraph_format.first_line_indent=Inches(-.25)
    p.paragraph_format.space_after=Pt(6)
    r=p.add_run(clean(text)); font(r)

def add_figure(doc,caption,path):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(6.1))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(9)
    r=c.add_run(clean(caption)); font(r,9,italic=True,color=GRAY)

def add_equation(doc,text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
    r=p.add_run(clean(text)); font(r,10.5,name="Cambria Math")

def parse(doc,path,stop_refs=False):
    lines=path.read_text(encoding="utf-8").splitlines(); i=1
    while i<len(lines):
        s=lines[i].strip()
        if stop_refs and s.lower()=="## references": break
        if not s or s.startswith("The acknowledgment wording follows"): i+=1; continue
        m=re.match(r"!\[(.*?)\]\((.*?)\)",s)
        if m: add_figure(doc,m.group(1),ROOT/m.group(2)); i+=1; continue
        if s=="\\[":
            eq=[]; i+=1
            while i<len(lines) and lines[i].strip()!="\\]": eq.append(lines[i].strip()); i+=1
            add_equation(doc," ".join(eq)); i+=1; continue
        if s.startswith("|"):
            block=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                cells=[x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?",x or "") for x in cells): block.append(cells)
                i+=1
            add_table(doc,block); continue
        if s.startswith("#"):
            level=len(s)-len(s.lstrip("#")); title=clean(s[level:].strip())
            if title:
                p=doc.add_paragraph(title,style="Heading "+str(min(3,max(1,level-1)))); p.paragraph_format.keep_with_next=True
            i+=1; continue
        mn=re.match(r"^\d+\.\s+(.*)",s)
        if mn: add_list(doc,mn.group(1),True); i+=1; continue
        if s.startswith("- "): add_list(doc,s[2:],False); i+=1; continue
        para(doc,s); i+=1

def configure(doc):
    sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492); sec.different_first_page_header_footer=True
    normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
    for name,size,color,before,after in [("Heading 1",16,BLUE,16,8),("Heading 2",13,BLUE,12,6),("Heading 3",12,DARK,8,4)]:
        st=doc.styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=color
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(hp.add_run("XPoSat POLIX | Faculty Guide Review Draft"),9,color=GRAY)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(fp.add_run("Faculty Guide Review Copy"),9,color=GRAY)

def cover(doc):
    for _ in range(5): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("FACULTY GUIDE REVIEW COPY"),11,bold=True,color=BLUE); p.paragraph_format.space_after=Pt(16)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("An Explainable AI Framework for Analysis of\nX-Ray Polarimetry Data from XPoSat"),24,bold=True,color=NAVY); p.paragraph_format.space_after=Pt(12)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("Combined Sample Research Paper Draft"),14,italic=True,color=GRAY); p.paragraph_format.space_after=Pt(28)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("Author One | Author Two | Author Three | Author Four"),11,bold=True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("Department of [Department Name]\n[College/University], [City, Country]"),10.5,color=GRAY); p.paragraph_format.space_after=Pt(34)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("Guide Review Brief | Decision Sheet | Manuscript Draft"),10.5,bold=True,color=DARK)

def part(doc,title,subtitle):
    doc.add_page_break(); p=doc.add_paragraph(title,style="Heading 1")
    s=doc.add_paragraph(); font(s.add_run(subtitle),11,italic=True,color=GRAY); s.paragraph_format.space_after=Pt(12)

def main():
    doc=Document(); configure(doc); cover(doc)
    part(doc,"Part I - Guide Review Brief","Concise summary for faculty discussion.")
    parse(doc,GUIDE)
    part(doc,"Part II - Guide Decision Sheet","Items requiring approval or domain review.")
    parse(doc,DECISIONS)
    part(doc,"Part III - Research Paper Draft","IEEE-conference-style content draft.")
    parse(doc,PAPER,True)
    cp=doc.core_properties; cp.title="An Explainable AI Framework for Analysis of X-Ray Polarimetry Data from XPoSat"; cp.author="Student Research Team"; cp.subject="Faculty guide review sample"
    doc.save(OUT); print(OUT)

if __name__=="__main__": main()
